"""Provisional V1-shaped attendance export from one native calendar snapshot."""
from __future__ import annotations

import calendar
import json
import re
from shutil import copyfileobj
from tempfile import SpooledTemporaryFile
from zipfile import ZIP_DEFLATED, ZipFile

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from grile.calendar_models import CalendarMonth, StoreHours
from services.exports.artifact import XlsxArtifact


def _workbook(data: CalendarMonth, site: str, codes: list[str]) -> Workbook:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = "Pontaj"
    hours = next((row for row in data.store_hours if row.site_code == site), StoreHours(site_code=site))
    sheet.merge_cells("A1:AH1")
    sheet['A1'] = "Foaie colectiva de prezenta — PROVIZORIU"
    sheet['B3'] = "Bază virtuală: TL" if site == "TL" else f"Magazin: {site}"
    sheet['A4'], sheet['B4'] = ("Tip", "Absențe · fără ore lucrate") if site == "TL" else ("Orar", f"{hours.opens}–{hours.closes}")
    sheet['B5'] = data.month
    workbook.properties.identifier = data.projection_revision
    for column, value in enumerate(["NrCrt", "Nume / Cod agent", *range(1, 32), "Total ore lucrate"], 1):
        sheet.cell(7, column, value)
    days = {(row.agent_code, row.work_date.day): row for row in data.attendance_days if row.site_code == site}
    totals = {row.agent_code: row.worked_minutes for row in data.attendance_by_store.get(site, [])}
    names = {r.agent_code: r.display_name for r in data.roster}
    for index in range(max(8, len(codes))):
        row = 8 + index * 3
        sheet.cell(row, 1, index + 1)
        sheet.cell(row + 1, 2, "Ora intrare-Ora iesire")
        sheet.cell(row + 2, 2, "Pauza")
        if index >= len(codes):
            continue
        code = codes[index]
        sheet.cell(row, 2, f'{names[code]} · {code}' if names.get(code) else code).data_type = 's'
        sheet.cell(row, 34, totals.get(code, 0) / 60)
        for number in range(1, 32):
            day = days.get((code, number))
            if day is None or day.status == 'off':
                continue
            sheet.cell(row, number + 2, 'CO' if day.status == 'leave' else day.worked_minutes / 60)
            if day.status == 'work':
                sheet.cell(row + 1, number + 2, f"{day.opens}–{day.closes}")
                sheet.cell(row + 2, number + 2, day.break_minutes / 60)
    _format(sheet, data.month, max(8, len(codes)))
    return workbook


def _format(sheet, month: str, count: int) -> None:
    year, number = map(int, month.split('-'))
    length = calendar.monthrange(year, number)[1]
    border = Border(*(Side(style='thin', color='999999') for _ in range(4)))
    for row in sheet.iter_rows(min_row=7, max_row=7 + 3 * count, max_col=34):
        for cell in row:
            cell.border = border
            cell.alignment = Alignment(horizontal='left' if cell.column == 2 else 'center', vertical='center')
            cell.font = Font(name='Calibri', size=10, bold=cell.row == 7)
            if 3 <= cell.column <= 33:
                day = cell.column - 2
                color = 'E5E7EB' if day > length else ('FFFF00' if calendar.weekday(year, number, day) >= 5 else 'FFFFFF')
                cell.fill = PatternFill('solid', fgColor=color)
    sheet.column_dimensions['A'].width = 6
    sheet.column_dimensions['B'].width = 26
    sheet.column_dimensions['AH'].width = 17
    for column in range(3, 34):
        sheet.column_dimensions[get_column_letter(column)].width = 13
    sheet.freeze_panes = 'C8'
    sheet.sheet_view.showGridLines = False
    sheet.page_setup.orientation = 'landscape'
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A3
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 1
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_area = f'A1:AH{7 + 3 * count}'
    sheet['A1'].font = Font(name='Calibri', size=14, bold=True)


def _attendance_manifest(data: CalendarMonth, sites: list[str]) -> str:
    return json.dumps({
        'month': data.month, 'projection_revision': data.projection_revision,
        'status': 'provisional', 'stores': [site for site in sites if site != 'TL'],
        'virtual_bases': ['TL'] if 'TL' in sites else [],
    }, ensure_ascii=False)


def build_attendance_zip(data: CalendarMonth) -> XlsxArtifact:
    active = [r for r in data.roster if r.active]
    participants = {r.agent_code for r in active} | {r.agent_code for r in data.attendance_days}
    sites = sorted({r.home_site_code for r in active if r.home_site_code != "TL"} | set(data.attendance_by_store) |
                   {r.site_code for r in data.store_hours})
    if not sites:
        raise HTTPException(409, "Confirm the monthly roster before exporting attendance")
    if len(sites) > 200 or len(participants) > 2000:
        raise HTTPException(422, "Attendance export exceeds the supported monthly cohort")
    stream = SpooledTemporaryFile(max_size=4 * 1024 * 1024, mode='w+b')
    try:
        with ZipFile(stream, 'w', compression=ZIP_DEFLATED) as archive:
            for index, site in enumerate(sites, 1):
                codes = sorted({r.agent_code for r in active if r.home_site_code == site} |
                               {r.agent_code for r in data.attendance_by_store.get(site, [])})
                if len(codes) > 100:
                    raise HTTPException(422, "Attendance export exceeds 100 participants per store")
                workbook = _workbook(data, site, codes)
                try:
                    with SpooledTemporaryFile(max_size=1024 * 1024, mode='w+b') as file:
                        workbook.save(file)
                        file.seek(0)
                        name = re.sub(r'[^\w.-]', '_', site)[:80]
                        with archive.open(f'{index:03d}_{name}_{data.month}.xlsx', 'w') as member:
                            copyfileobj(file, member, length=256 * 1024)
                finally:
                    workbook.close()
            archive.writestr('manifest.json', _attendance_manifest(data, sites))
        return XlsxArtifact(stream, f'Pontaje-provizorii-{data.month}.zip', stream.tell())
    except BaseException:
        stream.close()
        raise
