"""Provisional earnings and attendance from the same immutable read snapshot."""
from __future__ import annotations

import json
from shutil import copyfileobj
from tempfile import SpooledTemporaryFile
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

from grile.calendar_models import CalendarMonth
from grile.earnings_models import EarningsMonth
from services.grile_attendance_export import build_attendance_zip

_LABELS = {
    'salary_base': 'Bază salarială', 'vouchers': 'Bonuri', 'sim': 'SIM',
    'epay': 'E-pay', 'manual_incentives': 'Stimulente manuale',
    'confirmed': 'Confirmată', 'unavailable': 'Indisponibilă', 'conflicting': 'Contradictorie',
    'missing_published_source': 'Lipsește sursa publicată', 'missing_sales': 'Lipsesc vânzările',
    'missing_positive_target': 'Lipsește un target pozitiv', 'after_cutoff': 'Planificat după data limită',
}


def _rows(sheet, rows) -> None:
    for number, row in enumerate(rows, 1):
        sheet.append(row)
        for column, value in enumerate(row, 1):
            if isinstance(value, str):
                sheet.cell(number, column).data_type = 's'
    sheet.freeze_panes = 'C5'
    sheet.auto_filter.ref = f'A4:{sheet.cell(sheet.max_row, sheet.max_column).coordinate}'
    for cell in sheet[4]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(wrap_text=True)
    sheet.row_dimensions[4].height = 36
    for cells in sheet.columns:
        sheet.column_dimensions[cells[0].column_letter].width = 24


def _workbook(data: EarningsMonth) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    workbook.properties.identifier = data.projection_revision
    summary = workbook.create_sheet('Castiguri provizorii')
    _rows(summary, [
        ['PROVIZORIU — nu reprezintă salariul oficial', data.month],
        ['Vânzări până la', str(data.cutoff or 'Indisponibil')],
        ['Componente neincluse', ', '.join(_LABELS.get(value, value) for value in data.unavailable_components)],
        ['Cod agent', 'Nume confirmat', 'Magazin de bază', 'Identitate', 'Zile bază', 'Target personal (lei)',
         'Vânzări bază (lei)', 'Comision lunar (lei)', 'Comision alte locații (lei)', 'Plata suplimentărilor (lei)', 'Total calculat (lei)', 'Probleme'],
        *[[a.agent_code, a.display_name, a.home_site_code, _LABELS[a.identity_status], a.home_work_days,
           a.home_target, a.home_sales, a.home_commission, a.away_commission, a.supplemental_pay,
           a.known_earnings, ', '.join(_LABELS.get(value, value) for value in a.issues)] for a in data.agents],
    ])
    detail = workbook.create_sheet('Detalii zile')
    _rows(detail, [
        ['PROVIZORIU — atribuirea urmează calendarul confirmat', data.month],
        ['Vânzări până la', str(data.cutoff or 'Indisponibil')],
        ['Celule monetare goale = indisponibil; zilele viitoare rămân planificate'],
        ['Cod agent', 'Data', 'Magazin lucrat', 'Altă locație', 'Suplimentare', 'Vânzări',
         'Target zilnic', 'Comision alte locații', 'Plata suplimentării', 'Problemă'],
        *[[a.agent_code, d.work_date.isoformat(), d.site_code, d.away, d.supplemental,
           d.sales, d.daily_target, d.commission, d.supplemental_pay, _LABELS.get(d.issue, d.issue)]
          for a in data.agents for d in a.days],
    ])
    unassigned = workbook.create_sheet('Vanzari nealocate')
    _rows(unassigned, [
        ['PROVIZORIU — nu sunt atribuite automat unei persoane', data.month],
        ['Vânzări până la', str(data.cutoff or 'Indisponibil')],
        ['Calendar complet necesar pentru target; divizorul este numărul zilelor work programate'],
        ['Magazin', 'Data', 'Vânzări'],
        *[[r.site_code, r.sale_date.isoformat(), r.sales] for r in data.unassigned_sales],
    ])
    return workbook


def build_earnings_zip(calendar: CalendarMonth, earnings: EarningsMonth):
    if earnings.calendar_revision != calendar.projection_revision or earnings.month != calendar.month:
        raise HTTPException(409, 'Earnings and attendance must share one calendar revision')
    if len(earnings.agents) > 2000 or sum(len(a.days) for a in earnings.agents) > 62000 or len(earnings.unassigned_sales) > 62000:
        raise HTTPException(422, 'Earnings export exceeds the supported monthly cohort')
    artifact = build_attendance_zip(calendar)
    try:
        workbook = _workbook(earnings)
        try:
            with SpooledTemporaryFile(max_size=1024 * 1024, mode='w+b') as file:
                workbook.save(file)
                file.seek(0)
                with ZipFile(artifact.stream, 'a') as archive:
                    with archive.open(f'Castiguri-provizorii-{calendar.month}.xlsx', 'w') as member:
                        copyfileobj(file, member, length=256 * 1024)
                    archive.writestr('castiguri-manifest.json', json.dumps({
                        'month': earnings.month, 'status': 'provisional',
                        'projection_revision': earnings.projection_revision,
                        'calendar_revision': earnings.calendar_revision,
                        'cutoff': str(earnings.cutoff) if earnings.cutoff else None,
                        'selling_days': earnings.selling_days,
                        'unavailable_components': earnings.unavailable_components,
                    }, ensure_ascii=False))
        finally:
            workbook.close()
        artifact.filename = f'Grile-pontaje-provizorii-{calendar.month}.zip'
        artifact.size = artifact.stream.tell()
        return artifact
    except BaseException:
        artifact.close()
        raise
