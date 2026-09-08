"""The exported money and hours share one projection, with no duplicate pay."""
from io import BytesIO
import json
from unittest.mock import AsyncMock
from zipfile import ZipFile

from fastapi import HTTPException
from openpyxl import load_workbook
import pytest

from grile.earnings_projection import project_earnings
from services.grile_calendar import GrileCalendarService
from services.grile_earnings_export import build_earnings_zip
from test_grile_earnings import sources


def test_archive_reconciles_money_hours_revisions_and_literal_names():
    data = sources()
    data['calendar']['roster'][0].update(display_name='=Literal name', identity_status='confirmed')
    calendar = GrileCalendarService.project_calendar('2026-09', data['calendar'])
    earnings = project_earnings(calendar, data)
    artifact = build_earnings_zip(calendar, earnings)
    try:
        with ZipFile(artifact.stream) as archive:
            money_manifest = json.loads(archive.read('castiguri-manifest.json'))
            attendance_manifest = json.loads(archive.read('manifest.json'))
            assert money_manifest['calendar_revision'] == attendance_manifest['projection_revision']
            assert money_manifest['projection_revision'] == earnings.projection_revision
            workbook = load_workbook(BytesIO(archive.read('Castiguri-provizorii-2026-09.xlsx')), data_only=False)
            sheet = workbook['Castiguri provizorii']
            assert sheet['A5'].value == 'AG1'
            assert sheet['B5'].value == '=Literal name' and sheet['B5'].data_type == 's'
            assert sheet['K5'].value == 222
            assert sheet.max_row == 7  # one total per person, not one per worked store
            assert workbook.properties.identifier == earnings.projection_revision
            assert workbook['Detalii zile'].max_row == 10
            workbook.close()
            hours = []
            for name in archive.namelist():
                if name.endswith('.xlsx') and not name.startswith('Castiguri'):
                    book = load_workbook(BytesIO(archive.read(name)), data_only=True)
                    hours.append(book['Pontaj']['AH8'].value)
                    book.close()
            assert sum(hours) == 33  # same person: 22 home + 11 away
    finally:
        artifact.close()


@pytest.mark.asyncio
@pytest.mark.parametrize('change', ['sales', 'target', 'calendar', 'cutoff'])
async def test_export_rejects_changed_inputs_before_building(monkeypatch, change):
    data = sources()
    calendar = GrileCalendarService.project_calendar('2026-09', data['calendar'])
    revision = project_earnings(calendar, data).projection_revision
    if change == 'sales':
        data['sales'][0]['sales'] += 1
    elif change == 'target':
        data['targets'][0]['target_value'] += 1
    elif change == 'calendar':
        data['calendar']['days'][0]['revision'] += 1
    else:
        data['source'] = None
    read = AsyncMock(return_value=data)
    monkeypatch.setattr('services.grile_calendar.read_earnings_sources', read)
    with pytest.raises(HTTPException) as error:
        await GrileCalendarService(AsyncMock()).export_earnings('2026-09', revision)
    assert error.value.status_code == 409
    read.assert_awaited_once()


@pytest.mark.asyncio
async def test_export_missing_money_stays_blank_and_reads_once(monkeypatch):
    data = sources()
    data['source'] = None
    calendar = GrileCalendarService.project_calendar('2026-09', data['calendar'])
    revision = project_earnings(calendar, data).projection_revision
    read = AsyncMock(return_value=data)
    monkeypatch.setattr('services.grile_calendar.read_earnings_sources', read)
    artifact = await GrileCalendarService(AsyncMock()).export_earnings('2026-09', revision)
    try:
        with ZipFile(artifact.stream) as archive:
            book = load_workbook(BytesIO(archive.read('Castiguri-provizorii-2026-09.xlsx')), data_only=True)
            assert book['Castiguri provizorii']['K5'].value is None
            assert 'Bază salarială' in book['Castiguri provizorii']['A3'].value
            book.close()
    finally:
        artifact.close()
    read.assert_awaited_once()


def test_export_rejects_mixed_revisions_and_excessive_cohort():
    data = sources()
    calendar = GrileCalendarService.project_calendar('2026-09', data['calendar'])
    earnings = project_earnings(calendar, data)
    with pytest.raises(HTTPException) as error:
        build_earnings_zip(calendar, earnings.model_copy(update={'calendar_revision': 'wrong'}))
    assert error.value.status_code == 409
    with pytest.raises(HTTPException) as error:
        build_earnings_zip(calendar, earnings.model_copy(update={'agents': earnings.agents * 1000}))
    assert error.value.status_code == 422


def test_partial_build_failure_closes_artifact(monkeypatch):
    from services import grile_earnings_export as module
    data = sources()
    calendar = GrileCalendarService.project_calendar('2026-09', data['calendar'])
    earnings = project_earnings(calendar, data)
    artifact = module.build_attendance_zip(calendar)
    monkeypatch.setattr(module, 'build_attendance_zip', lambda _: artifact)
    def fail(*_):
        raise OSError('synthetic disk failure')
    monkeypatch.setattr(module, '_workbook', fail)
    with pytest.raises(OSError):
        module.build_earnings_zip(calendar, earnings)
    assert artifact.stream.closed


def test_full_calendar_distinguishes_unallocated_leave_off_cancelled_and_future_work():
    from datetime import date
    data = sources()
    for number, status in [(4, 'leave'), (5, 'off'), (6, 'cancelled')]:
        data['calendar']['days'].append(dict(agent_code='AG1', site_code='A', work_date=date(2026, 9, number),
                                            status=status, supplemental=False, revision=1))
    data['calendar']['days'].append(dict(agent_code='SUP', site_code='B', work_date=date(2026, 9, 8),
                                        status='work', supplemental=False, revision=1))
    calendar = GrileCalendarService.project_calendar('2026-09', data['calendar'])
    earnings = project_earnings(calendar, data)
    artifact = build_earnings_zip(calendar, earnings)
    try:
        with ZipFile(artifact.stream) as archive:
            book = load_workbook(BytesIO(archive.read('Castiguri-provizorii-2026-09.xlsx')), data_only=True)
            summary = book['Castiguri provizorii']
            assert summary['M5'].value == 33 and summary['N5'].value == 1
            rows = list(book['Calendar'].iter_rows(min_row=5, values_only=True))
            assert len(rows) == 63
            assert sum(row[4] == 'Nealocat' for row in rows) == 53
            actual = {row[1]: row for row in rows if row[2] == 'AG1' and row[0] == 'A'}
            assert [actual[f'2026-09-0{n}'][4] for n in (4, 5, 6)] == ['Concediu', 'Liber', 'Anulat']
            assert (actual['2026-09-04'][10], actual['2026-09-05'][10], actual['2026-09-06'][10]) == (0, 0, 0)
            assert sum(row[10] or 0 for row in rows) == 77  # seven programmed days at 11 hours
            future = next(row for row in book['Detalii zile'].iter_rows(min_row=5, values_only=True) if row[1] == '2026-09-08')
            assert future[5] is None and future[9] == 'Planificat după data limită'
            book.close()
    finally:
        artifact.close()


def test_leap_month_keeps_unallocated_days_and_inactive_roster_out_of_cohort():
    data = sources()
    data['calendar']['days'] = []
    for entry in data['calendar']['roster']:
        entry['month'] = '2028-02'
    data['calendar']['roster'][2].update(active=False, home_site_code='INACTIVE')
    calendar = GrileCalendarService.project_calendar('2028-02', data['calendar'])
    data.update(sales=[], targets=[], source=None)
    earnings = project_earnings(calendar, data)
    artifact = build_earnings_zip(calendar, earnings)
    try:
        with ZipFile(artifact.stream) as archive:
            book = load_workbook(BytesIO(archive.read('Castiguri-provizorii-2028-02.xlsx')), data_only=True)
            rows = list(book['Calendar'].iter_rows(min_row=5, values_only=True))
            assert len(rows) == 29 and rows[-1][1] == '2028-02-29'
            assert all(row[0] == 'A' and row[4] == 'Nealocat' for row in rows)
            book.close()
    finally:
        artifact.close()


def test_calendar_keeps_cancelled_only_location_without_adding_paid_hours():
    from datetime import date
    data = sources()
    data['calendar']['days'].append(dict(
        agent_code='AG1', site_code='CANCELLED_ONLY', work_date=date(2026, 9, 6),
        status='cancelled', supplemental=False, revision=1,
    ))
    calendar = GrileCalendarService.project_calendar('2026-09', data['calendar'])
    artifact = build_earnings_zip(calendar, project_earnings(calendar, data))
    try:
        with ZipFile(artifact.stream) as archive:
            book = load_workbook(BytesIO(archive.read('Castiguri-provizorii-2026-09.xlsx')))
            rows = [row for row in book['Calendar'].iter_rows(min_row=5, values_only=True)
                    if row[0] == 'CANCELLED_ONLY']
            assert len(rows) == 31  # 30 unallocated days plus the cancellation
            cancelled = next(row for row in rows if row[4] == 'Anulat')
            assert cancelled[1:3] == ('2026-09-06', 'AG1')
            assert cancelled[10] == 0
            assert book['Castiguri provizorii']['K5'].value == 222
            book.close()
    finally:
        artifact.close()
