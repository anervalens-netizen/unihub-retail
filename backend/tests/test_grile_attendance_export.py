from datetime import date
from io import BytesIO
import json
from unittest.mock import AsyncMock
from zipfile import ZipFile

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook
from pydantic import ValidationError

from grile.calendar_models import StoreHoursInput
from services.grile_calendar import GrileCalendarService


def repository():
    return AsyncMock(read=AsyncMock(return_value={
        'roster': [dict(month='2026-09', agent_code='AG1', home_site_code='A', active=True, revision=1)],
        'days': [dict(work_date=date(2026, 9, i), agent_code='AG1', site_code=site,
                      status=status, supplemental=site == 'B' and status == 'work', revision=1)
                 for i, site, status in [(1, 'A', 'work'), (2, 'B', 'work'), (3, 'A', 'leave'),
                                         (4, 'A', 'off'), (5, 'B', 'cancelled')]],
        'store_hours': [dict(site_code='B', opens='09:00', closes='22:00', break_minutes=60, revision=1)],
    }))


@pytest.mark.asyncio
async def test_zip_reconciles_actual_store_minutes_and_reference_layout():
    service = GrileCalendarService(repository())
    data = await service.read('2026-09')
    assert data.attendance[0].worked_minutes == 23 * 60
    artifact = await service.export_attendance('2026-09', data.projection_revision)
    try:
        with ZipFile(artifact.stream) as archive:
            manifest = json.loads(archive.read('manifest.json'))
            assert manifest['projection_revision'] == data.projection_revision
            files = [name for name in archive.namelist() if name.endswith('.xlsx')]
            assert len(files) == 2
            first = load_workbook(BytesIO(archive.read(files[0])), data_only=True)
            second = load_workbook(BytesIO(archive.read(files[1])), data_only=True)
            a, b = first['Pontaj'], second['Pontaj']
            assert (a['C8'].value, a['C9'].value, a['C10'].value) == (11, '10:00–22:00', 1)
            assert (b['D8'].value, b['D9'].value, b['D10'].value) == (12, '09:00–22:00', 1)
            assert a['D8'].value is None  # no duplicate at home
            assert a['E8'].value == 'CO' and a['E9'].value is None
            assert a['F8'].value is None and b['G8'].value is None
            assert a['AH8'].value + b['AH8'].value == 23
            assert a['G8'].fill.fgColor.rgb.endswith('FFFF00')  # Saturday 5 Sept
            assert a['AG8'].fill.fgColor.rgb.endswith('E5E7EB')  # day 31 absent
            assert a['B31'].value == 'Pauza'
            assert a.freeze_panes == 'C8'
            first.close()
            second.close()
    finally:
        artifact.close()


@pytest.mark.asyncio
async def test_changed_hours_invalidate_export_revision():
    repo = repository()
    service = GrileCalendarService(repo)
    first = await service.read('2026-09')
    repo.read.return_value['store_hours'][0]['opens'] = '10:00'
    with pytest.raises(HTTPException) as error:
        await service.export_attendance('2026-09', first.projection_revision)
    assert error.value.status_code == 409


@pytest.mark.parametrize('values', [dict(opens='22:00', closes='09:00'), dict(break_minutes=720),
                                    dict(opens='24:00'), dict(break_minutes=-1)])
def test_invalid_store_intervals_rejected(values):
    with pytest.raises(ValidationError):
        StoreHoursInput(expected_revision=0, **values)
